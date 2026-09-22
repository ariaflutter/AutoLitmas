#!/usr/bin/env python3
"""
write_excel.py — Tulis JSON Litmas (146 field) ke file .xlsx sesuai struktur Contoh Fix.ods.

Usage:
    python3 write_excel.py --json output_klien.json --out "Litmas_Mohammad_Zaenal_Abidin.xlsx"
    python3 write_excel.py --json output_klien.json --out "Litmas_X.xlsx" --append "Contoh Fix.xlsx"

Jika --append diberikan, data ditambahkan sebagai row baru di file target (header row 1).
Jika tidak, file baru dibuat dengan header row 1 + data row 2.

Struktur output = persis Contoh Fix.ods: 146 kolom, header row 1, data mulai row 2.
"""
import argparse
import json
import os
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl tidak terinstall.", file=sys.stderr)
    print("Install dengan: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_HEADERS_FILE = SCRIPT_DIR.parent / "templates" / "reference_contoh_fix.json"

FIELD_ORDER = [
    "1_Nama_Klien", "2_Status_Litmas", "3_Jenis_Litmas", "4_Tanggal_TPP",
    "5_Nomor_Sidang_TPP", "6_Tanggal_Surat_Pengantar_Laporan",
    "7_Perihal_Surat_Pengantar_Laporan_Litmas", "8_Peminta_Nama_Instansi",
    "9_Peminta_Kelas_Instansi", "10_Peminta_Kota_Instansi",
    "11_Tanggal_Permintaan_Surat", "12_Nomor_Permintaan_Surat",
    "13_Peminta_Perihal_Surat", "14_No_Surat_Tugas", "15_Tanggal_Surat_Tugas",
    "16_No_Register_Litmas", "17_No_Register_Lapas", "18_Tempat_Lahir",
    "19_Tanggal_Lahir", "20_Tanggal_Putusan", "21_No_Putusan", "22_Pidana",
    "23_Pertama_Ditahan", "24_Sepertiga", "25_Setengah", "26_Dua_Pertiga",
    "27_Ekspirasi", "28_Usia", "29_Perkara", "30_Pasal", "31_Alamat",
    "32_Jenis_Kelamin", "33_Agama", "34_Suku", "35_Bangsa", "36_Warga_Negara",
    "37_Pendidikan_Terakhir", "38_Pekerjaan", "39_Status_Perkawinan",
    "40_Ciri_Khusus", "41_Ayah_Nama", "42_Ayah_Tempat_Lahir",
    "43_Ayah_Tanggal_Lahir", "44_Ayah_Agama", "45_Ayah_Suku",
    "46_Ayah_Bangsa", "47_Ayah_Warga_Negara", "48_Ayah_Pendidikan_Terakhir",
    "49_Ayah_Pekerjaan", "50_Ayah_Alamat", "51_Ayah_Hubungan", "52_Ibu_Nama",
    "53_Ibu_Tempat_Lahir", "54_Ibu_Tanggal_Lahir", "55_Ibu_Agama",
    "56_Ibu_Suku", "57_Ibu_Bangsa", "58_Ibu_Warga_Negara",
    "59_Ibu_Pendidikan_Terakhir", "60_Ibu_Pekerjaan", "61_Ibu_Alamat",
    "62_Ibu_Hubungan", "63_Istri_Nama", "64_Istri_Tempat_Lahir",
    "65_Istri_Tanggal_Lahir", "66_Istri_Agama", "67_Istri_Suku",
    "68_Istri_Bangsa", "69_Istri_Warga_Negara", "70_Istri_Pendidikan_Terakhir",
    "71_Istri_Pekerjaan", "72_Istri_Alamat", "73_Istri_Hubungan",
    "74_Penjamin_Nama", "75_Penjamin_Tempat_Lahir", "76_Penjamin_Tanggal_Lahir",
    "77_Penjamin_Agama", "78_Penjamin_Suku", "79_Penjamin_Bangsa",
    "80_Penjamin_Warga_Negara", "81_Penjamin_Pendidikan_Terakhir",
    "82_Penjamin_Pekerjaan", "83_Penjamin_Alamat", "84_Penjamin_Hubungan",
    "85_Riwayat_Kelahiran_Klien", "86_Riwayat_Pertumbuhan_Fisik_Klien",
    "87_Riwayat_Perkembangan_Psiko_Sosial_Klien", "88_Pendidikan_dalam_Keluarga",
    "89_Pendidikan_Formal", "90_Pendidikan_Non_Formal",
    "91_Bakat_dan_Potensi_Klien",
    "92_Relasi_Sosial_dengan_Orang_tua_dan_Keluarga",
    "93_Ketaatan_Klien_dalam_Beragama", "94_Kebiasaan_Klien_yang_Baik",
    "95_Kebiasaan_Klien_yang_Buruk", "96_Sikap_Klien_dalam_Bekerja",
    "97_Riwayat_Pelanggaran_Hukum",
    "98_Riwayat_Mengkonsumsi_Rokok_Napza_dan_Alkohol",
    "99_Riwayat_Perkawinan_Klien", "100_Riwayat_Perkawinan_Orang_tua",
    "101_Relasi_Sosial_dalam_keluarga",
    "102_Pekerjaan_dan_Keadaan_ekonomi",
    "103_Keadaan_Rumah_dan_Tempat_Tinggal",
    "104_Relasi_Sosial_dengan_Masyarakat",
    "105_Kondisi_Sosial_Alam_tempat_Orang_tua",
    "106_Penggolongan_Profesi_dan_Mata_Pencaharian",
    "107_Stratifikasi_Sosial_Ekonomi_masyarakat",
    "108_Tingkat_Pendidikan_Rata_Rata_Masyarakat",
    "109_Pola_Hubungan_Interaksi_Sosial_dalam_Masyarakat",
    "110_Kegiatan_Pendidikan",
    "111_Kepedulian_terhadap_Kegiatan_Keagamaan",
    "112_Kesadaran_terhadap_Kepatuhan_Nilai_Norma_Hukum_yang_Berlaku",
    "113_Riwayat_Perkawinan_Penjamin",
    "114_Relasi_Sosial_dalam_keluarga_Penjamin",
    "115_Relasi_Sosial_dalam_Masyarakat_Penjamin",
    "116_Pekerjaan_dan_Keadaan_ekonomi_Penjamin",
    "117_Keadaan_Rumah_dan_Tempat_Tinggal_Penjamin",
    "118_Kondisi_Sosial_Alam_tempat_penjamin",
    "119_Latar_Belakang_Tindak_Pidana", "120_Kronologi",
    "121_Keadaan_Korban", "122_Akibat_yang_ditimbulkan_Kepada_Korban",
    "123_Akibat_Tindak_Pidana_Terhadap_Klien",
    "124_Akibat_Tindak_Pidana_Terhadap_Orang_tua",
    "125_Akibat_Tindak_Pidana_Terhadap_Masyarakat",
    "126_Sikap_dan_Tanggapan_Klien",
    "127_Sikap_dan_Tanggapan_Orang_tua_Keluarga",
    "128_Sikap_dan_Tanggapan_Korban",
    "129_Sikap_dan_Tanggapan_Masyarakat",
    "130_Sikap_dan_Tanggapan_Pemerintah_Setempat",
    "131_Evaluasi_Pelaksanaan_Program_Admisi_Orientasi_dan_Observasi",
    "132_Program_Pembinaan_Kepribadian",
    "133_Program_Pembinaan_Kemandirian",
    "134_Relasi_Sosial_selama_di_dalam_Lapas_ke_Sesama_WBP",
    "135_Relasi_Sosial_Selama_di_dalam_Lapas_kepada_Petugas",
    "136_Relasi_Sosial_selama_di_dalam_Lapas_Kepada_Keluarga",
    "137_Relasi_Sosial_di_dalam_Lapas_kepada_Masyarakat_Luar",
    "138_Evaluasi_Program_Asimilasi",
    "139_Asesmen_RRI_dan_Kriminogenik",
    "140_Sikap_Klien_Menjalani_Masa_Pembinaan_dan_Risiko_Pengulangan_Tindak_Pidana",
    "141_Perkembangan_dan_Perubahan_Perilaku_Klien_setelah_mengikuti_program_pembinaan",
    "142_Penerimaan_dan_kesiapan_masyarakat_pemerintah_dan_korban",
    "143_Hasil_Evaluasi_Program_Asimilasi",
    "144_Kemungkinan_Memenuhi_Syarat_atau_tidak_untuk_Program_Re_integrasi",
    "145_Kesimpulan", "146_Rekomendasi",
]

assert len(FIELD_ORDER) == 146, f"Expected 146 fields, got {len(FIELD_ORDER)}"


def load_headers():
    """Load 146 headers dari reference_contoh_fix.json."""
    with open(DEFAULT_HEADERS_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
    headers = data["headers"]
    assert len(headers) == 146, f"Expected 146 headers, got {len(headers)}"
    return headers


def normalize_value(v):
    """Normalize value: None -> '', keep string as-is."""
    if v is None:
        return ""
    return str(v)


def build_row(data):
    """Build ordered list of 146 values from JSON data dict (key = field name)."""
    row = []
    missing = []
    for field in FIELD_ORDER:
        if field in data:
            row.append(normalize_value(data[field]))
        else:
            # try numeric prefix match (e.g. "1" -> "1_Nama_Klien")
            num = field.split("_")[0]
            if num in data:
                row.append(normalize_value(data[num]))
            else:
                row.append("")
                missing.append(field)
    if missing:
        print(f"WARN: {len(missing)} field kosong: {missing[:5]}...", file=sys.stderr)
    return row


def style_header(ws, n_cols=146):
    """Apply header styling: bold, green bg, border, wrap, center."""
    header_font = Font(name="Arial", size=12, bold=True, color="000000")
    header_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="middle", wrap_text=True)
    thin = Side(border_style="thin", color="000000")
    header_border = Border(left=thin, right=thin, bottom=thin)

    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    ws.row_dimensions[1].height = 40


def style_data_row(ws, row_idx, n_cols=146):
    """Apply data row styling: wrap text, top align."""
    data_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.alignment = data_align


def set_column_widths(ws, n_cols=146):
    """Set reasonable column widths. Kolom 1-84 (terstruktur) lebih sempit, 85-146 (narasi) lebar."""
    for col_idx in range(1, n_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx <= 84:
            ws.column_dimensions[col_letter].width = 20
        else:
            ws.column_dimensions[col_letter].width = 60


def write_new(data, out_path):
    """Buat file Excel baru: header row 1 + data row 2."""
    headers = load_headers()
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Header row 1
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # Data row 2
    row = build_row(data)
    for col_idx, value in enumerate(row, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    style_header(ws)
    style_data_row(ws, 2)
    set_column_widths(ws)

    wb.save(out_path)
    print(f"OK: {out_path} (146 kolom, 1 row data)")
    return out_path


def write_append(data, target_path):
    """Append data sebagai row baru di file target (header row 1)."""
    if not os.path.exists(target_path):
        print(f"ERROR: file target tidak ada: {target_path}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(target_path)
    ws = wb.active

    # Cari row kosong berikutnya
    next_row = ws.max_row + 1
    # skip kalau row terakhir kosong
    while next_row > 2 and all(
        (ws.cell(row=next_row - 1, column=c).value in (None, ""))
        for c in range(1, 147)
    ):
        next_row -= 1

    row = build_row(data)
    for col_idx, value in enumerate(row, start=1):
        ws.cell(row=next_row, column=col_idx, value=value)

    style_data_row(ws, next_row)
    wb.save(target_path)
    print(f"OK: appended row {next_row} ke {target_path}")
    return target_path


def main():
    parser = argparse.ArgumentParser(
        description="Tulis JSON Litmas (146 field) ke file .xlsx"
    )
    parser.add_argument(
        "--json", required=True, help="Path file JSON output klien (146 field)"
    )
    parser.add_argument(
        "--out", required=True, help="Path file .xlsx output (file baru)"
    )
    parser.add_argument(
        "--append",
        help="Path file .xlsx target untuk append row baru (opsional)",
        default=None,
    )
    args = parser.parse_args()

    with open(args.json, "r", encoding="utf-8") as f:
        data = json.load(f)

    if args.append:
        write_append(data, args.append)
    else:
        write_new(data, args.out)


if __name__ == "__main__":
    main()
